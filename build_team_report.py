from pathlib import Path

import pandas as pd
from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


PROJECT = Path(__file__).resolve().parent
RESULTS = PROJECT / "results"
ASSETS = PROJECT / "report_assets_v3"
OUTPUT = PROJECT / "report" / "팀프로젝트_보고서_사용자재연결예측_보수적최종본.docx"

NAVY = "17365D"
PALE_BLUE = "EAF2F8"
PALE_GRAY = "F4F6F8"
MID_GRAY = "667085"
LIGHT_BORDER = "D9D9D9"
BLACK = "000000"
WHITE = "FFFFFF"


def set_run_font(run, name="맑은 고딕", size=10.8, bold=None, color=BLACK):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run._element.rPr.rFonts.set(qn("w:ascii"), name)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), name)
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    if bold is not None:
        run.bold = bold


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_border(cell, color=LIGHT_BORDER, size="5"):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = borders.find(qn("w:" + edge))
        if el is None:
            el = OxmlElement("w:" + edge)
            borders.append(el)
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), size)
        el.set(qn("w:color"), color)


def set_cell_width(cell, width_cm):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(int(Cm(width_cm).emu / 635)))
    tc_w.set(qn("w:type"), "dxa")


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    flag = OxmlElement("w:tblHeader")
    flag.set(qn("w:val"), "true")
    tr_pr.append(flag)


def prevent_row_split(row):
    tr_pr = row._tr.get_or_add_trPr()
    tr_pr.append(OxmlElement("w:cantSplit"))


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    sep = OxmlElement("w:fldChar")
    sep.set(qn("w:fldCharType"), "separate")
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, sep, end])
    set_run_font(run, size=9, color=MID_GRAY)


def add_text(doc, text, size=10.8, bold=False, first_indent=True, spacing=1.45):
    p = doc.add_paragraph()
    p.paragraph_format.line_spacing = spacing
    p.paragraph_format.space_after = Pt(6)
    if first_indent:
        p.paragraph_format.first_line_indent = Cm(0.72)
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    r = p.add_run(text)
    set_run_font(r, size=size, bold=bold)
    return p


def add_lead(doc, lead, text):
    p = doc.add_paragraph()
    p.paragraph_format.line_spacing = 1.35
    p.paragraph_format.space_after = Pt(6)
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    r1 = p.add_run(lead + "  ")
    set_run_font(r1, bold=True)
    r2 = p.add_run(text)
    set_run_font(r2)
    return p


def add_bullets(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.left_indent = Cm(0.72)
        p.paragraph_format.first_line_indent = Cm(-0.34)
        p.paragraph_format.line_spacing = 1.3
        p.paragraph_format.space_after = Pt(4)
        r = p.add_run(item)
        set_run_font(r)


def add_heading(doc, text, level=1):
    p = doc.add_paragraph(style=f"Heading {level}")
    p.paragraph_format.keep_with_next = True
    # Give subsection headings an explicit line box and enough separation from
    # the following paragraph.  LibreOffice otherwise compresses a heading and
    # the next list/paragraph when the page is nearly full.
    p.paragraph_format.line_spacing = 1.2
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(8)
    r = p.add_run(text)
    set_run_font(r, size=16 if level == 1 else 12.5, bold=True, color=BLACK)
    return p


def add_table(doc, headers, rows, widths, font_size=9.0):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.style = "Table Grid"
    header = table.rows[0]
    set_repeat_table_header(header)
    prevent_row_split(header)
    for idx, value in enumerate(headers):
        cell = header.cells[idx]
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cell_width(cell, widths[idx])
        set_cell_shading(cell, NAVY)
        set_cell_border(cell)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1.1
        r = p.add_run(str(value))
        set_run_font(r, size=font_size, bold=True, color=WHITE)
    for ridx, values in enumerate(rows):
        row = table.add_row()
        prevent_row_split(row)
        for cidx, value in enumerate(values):
            cell = row.cells[cidx]
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_width(cell, widths[cidx])
            set_cell_border(cell)
            if ridx % 2 == 1:
                set_cell_shading(cell, PALE_GRAY)
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT if cidx in (0, len(values) - 1) else WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.08
            r = p.add_run(str(value))
            set_run_font(r, size=font_size)
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(2)
    return table


def add_caption(doc, text, kind="그림"):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(9)
    p.paragraph_format.keep_together = True
    r = p.add_run(f"{kind} {text}")
    set_run_font(r, size=9.2, color=MID_GRAY)


def add_figure(doc, path, caption, width=6.0):
    path = Path(path)
    if not path.exists():
        raise FileNotFoundError(path)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.space_before = Pt(5)
    p.paragraph_format.space_after = Pt(0)
    p.add_run().add_picture(str(path), width=Inches(width))
    add_caption(doc, caption)


def add_flow(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(9)
    p.paragraph_format.space_after = Pt(10)
    p.paragraph_format.keep_together = True
    r = p.add_run(text)
    set_run_font(r, size=11.2, bold=True, color=NAVY)


def build_primary_ap_figure(dataset, title, prevalence, output):
    """Plot only the models retained in the primary descriptive comparison."""
    table = pd.read_csv(RESULTS / "research_v2" / dataset / "model_summary.csv").set_index("model")
    models = ["Full LightGBM", "TGN recurrence", "History HGB", "Rolling GraphSAGE", "Pair frequency"]
    labels = ["Full LightGBM", "TGN 재발생", "History HGB", "Rolling GraphSAGE", "쌍 빈도 기준선"]
    values = table.loc[models, "ap_mean"].astype(float).to_numpy()
    errors = table.loc[models, "ap_fold_sd"].astype(float).to_numpy()
    canvas = Image.new("RGB", (1800, 1000), "white")
    draw = ImageDraw.Draw(canvas)
    regular = ImageFont.truetype(r"C:\Windows\Fonts\malgun.ttf", 34)
    small = ImageFont.truetype(r"C:\Windows\Fonts\malgun.ttf", 27)
    bold = ImageFont.truetype(r"C:\Windows\Fonts\malgunbd.ttf", 48)
    left, right, top, bottom = 430, 1650, 180, 830
    maximum = max(float((values + errors).max()) + 0.03, 0.27)

    def x(value):
        return left + int((right - left) * float(value) / maximum)

    title_box = draw.textbbox((0, 0), title, font=bold)
    draw.text(((1800 - (title_box[2] - title_box[0])) / 2, 58), title, fill="#111111", font=bold)
    for tick in (0.00, 0.05, 0.10, 0.15, 0.20, 0.25):
        xx = x(tick)
        draw.line((xx, top - 25, xx, bottom), fill="#D9E0E5", width=2)
        draw.text((xx - 28, bottom + 20), f"{tick:.2f}", fill="#667085", font=small)
    px = x(prevalence)
    for yy in range(top - 35, bottom, 18):
        draw.line((px, yy, px, min(yy + 10, bottom)), fill="#777777", width=4)
    draw.text((px + 12, top - 65), f"양성률 {prevalence:.4f}", fill="#666666", font=small)
    for idx, (label, value, error) in enumerate(zip(labels, values, errors)):
        cy = top + idx * 125
        draw.text((35, cy + 12), label, fill="#222222", font=regular)
        draw.rounded_rectangle((left, cy, x(value), cy + 62), radius=6, fill="#24657A")
        low, high = x(max(value - error, 0)), x(value + error)
        draw.line((low, cy + 31, high, cy + 31), fill="#111111", width=5)
        draw.line((low, cy + 20, low, cy + 42), fill="#111111", width=5)
        draw.line((high, cy + 20, high, cy + 42), fill="#111111", width=5)
        # Put the value beyond the upper error-bar cap so the label remains readable.
        draw.text((high + 16, cy + 12), f"{value:.4f}", fill="#222222", font=small)
    note = "5개 시간 폴드 평균 AP  오차막대는 폴드 표준편차"
    note_box = draw.textbbox((0, 0), note, font=small)
    draw.text(((1800 - (note_box[2] - note_box[0])) / 2, 920), note, fill="#667085", font=small)
    output.parent.mkdir(parents=True, exist_ok=True)
    canvas.save(output)


PRIMARY_MO_FIGURE = ASSETS / "mathoverflow_ap_primary_ko.png"
PRIMARY_AU_FIGURE = ASSETS / "askubuntu_ap_primary_ko.png"
build_primary_ap_figure("mathoverflow", "MathOverflow 주 모델 비교", 0.1210, PRIMARY_MO_FIGURE)
build_primary_ap_figure("askubuntu", "AskUbuntu 주 모델 비교", 0.0722, PRIMARY_AU_FIGURE)


doc = Document()
section = doc.sections[0]
section.page_width = Inches(8.5)
section.page_height = Inches(11)
section.top_margin = Inches(0.78)
section.bottom_margin = Inches(0.72)
section.left_margin = Inches(0.85)
section.right_margin = Inches(0.85)
section.header_distance = Inches(0.3)
section.footer_distance = Inches(0.35)
section.different_first_page_header_footer = True

styles = doc.styles
normal = styles["Normal"]
normal.font.name = "맑은 고딕"
normal._element.rPr.rFonts.set(qn("w:eastAsia"), "맑은 고딕")
normal.font.size = Pt(10.8)
normal.font.color.rgb = RGBColor.from_string(BLACK)

title_style = styles["Title"]
title_style.font.name = "맑은 고딕"
title_style._element.rPr.rFonts.set(qn("w:eastAsia"), "맑은 고딕")
title_style.font.size = Pt(28)
title_style.font.bold = True
title_style.font.color.rgb = RGBColor.from_string(BLACK)
title_ppr = title_style._element.get_or_add_pPr()
title_border = title_ppr.find(qn("w:pBdr"))
if title_border is not None:
    title_ppr.remove(title_border)

for level, size in ((1, 16), (2, 12.5), (3, 11.2)):
    style = styles[f"Heading {level}"]
    style.font.name = "맑은 고딕"
    style._element.rPr.rFonts.set(qn("w:eastAsia"), "맑은 고딕")
    style.font.size = Pt(size)
    style.font.bold = True
    style.font.color.rgb = RGBColor.from_string(BLACK)

doc.core_properties.title = "시간 네트워크 기반 사용자 재연결 예측"
doc.core_properties.subject = "2개월 학부 팀프로젝트 보고서"
doc.core_properties.keywords = "시간 네트워크 사용자 재연결 링크 예측 LightGBM GraphSAGE TGN"

for sec in doc.sections:
    add_page_number(sec.footer.paragraphs[0])


# 표지
p = doc.add_paragraph()
p.paragraph_format.space_before = Pt(42)
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("학부 팀프로젝트 보고서")
set_run_font(r, size=14, bold=True, color=MID_GRAY)

p = doc.add_paragraph(style="Title")
p.paragraph_format.space_before = Pt(68)
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("직전 4시간 상호작용 사용자쌍의\n다음 24시간 재발생 순위화")
set_run_font(r, size=28, bold=True, color=BLACK)

p = doc.add_paragraph()
p.paragraph_format.space_before = Pt(22)
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("MathOverflow AskUbuntu Super User 데이터 분석")
set_run_font(r, size=14, color=BLACK)

p = doc.add_paragraph()
p.paragraph_format.space_before = Pt(120)
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("수행 기간 2개월")
set_run_font(r, size=12, bold=True, color=MID_GRAY)

p = doc.add_paragraph()
p.paragraph_format.space_before = Pt(12)
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("2026년 10월")
set_run_font(r, size=11, color=MID_GRAY)

doc.add_page_break()


# 한눈에 보기
add_heading(doc, "보고서 요약", 1)
add_text(doc, "이 프로젝트의 주 실험은 온라인 커뮤니티에서 예측 직전 4시간에 이미 상호작용한 방향성 사용자쌍만 후보로 삼고, 그 쌍이 다음 24시간 안에 다시 관측되는지를 순위화한다. 모든 사용자 사이의 신규 링크를 찾는 일반적인 링크 예측이 아니라 조건부 단기 재발생 예측이다.")
add_text(doc, "프로젝트는 인터넷 밈 확산과 X 리포스트 네트워크 분석에서 시작했다. X API 비용과 수집 재현성 문제로 공개 데이터로 전환했고, Higgs Twitter에서 시간 예측 구조를 시험한 뒤 기간이 더 긴 MathOverflow와 AskUbuntu로 분석을 확장했다. MathOverflow는 후보 기간과 특성 구성을 탐색한 회고적 실험이고, AskUbuntu는 고정한 절차를 해당 자료에서 다시 학습한 도메인 반복 실험이다. 완전히 독립적인 제로샷 검증으로 보지 않는다. Super User는 별도의 사전 고정 절차를 사용했지만 후보 기간이 24시간이므로 4시간 결과와 직접 비교하지 않는다.")
add_text(doc, "Full LightGBM의 5개 시간 폴드 평균 AP는 MathOverflow 0.2260, AskUbuntu 0.2183이었다. 양성률은 각각 0.1210과 0.0722였고 P@5는 0.2292와 0.2608이었다. 시드 점수를 먼저 평균한 뒤 적용한 시간 블록 부트스트랩에서 Full LightGBM과 History HGB의 pooled AP 차이는 7일, 14일, 28일 블록 모두 양의 구간을 보였다. 이 구간은 선택한 재표집 가정 아래의 불확실성만 나타내며 학습 시드 변동 전체를 포함하지 않는다. 결과는 과거 정보에 순위 신호가 있었음을 보여주지만 신규 링크 예측, 플랫폼 간 전이, 모델의 보편적 우월성 또는 서비스 효과를 입증하지 않는다.")

add_table(
    doc,
    ["항목", "내용"],
    [
        ["주 질문", "직전 4시간에 관측된 방향성 사용자쌍이 다음 24시간에 다시 관측되는가"],
        ["탐색 데이터", "MathOverflow"],
        ["반복 데이터", "AskUbuntu 도메인별 재학습"],
        ["추가 확인", "Super User 24시간 후보 별도 실험"],
        ["주 지표", "평균 정밀도 AP와 상위 5개 후보의 P@5 R@5"],
        ["핵심 결론", "선택한 조건에서 관계 이력과 행동 유형에 순위 신호가 관찰됨"],
    ],
    widths=[4.1, 12.1],
    font_size=9.4,
)

doc.add_page_break()


# 1장
add_heading(doc, "1 프로젝트 개요", 1)
add_heading(doc, "1.1 예측하려는 것", 2)
add_text(doc, "온라인 커뮤니티에서는 사용자가 질문에 답하거나 댓글을 남기면서 다른 사용자와 상호작용한다. 이 프로젝트는 최근에 관측된 방향성 사용자쌍 안에서 다음 24시간 재발생 후보를 순위화한다. 추정 대상은 P(다음 24시간 재발생 | 직전 4시간 상호작용)이며, 한 번도 연결되지 않은 쌍과 오래전에만 연결된 쌍은 주 평가 대상이 아니다. 결과는 연구용 순위로 해석하며 서비스 배포 효과까지 주장하지 않는다.")
add_flow(doc, "과거 상호작용  →  후보 사용자 쌍  →  모델 점수  →  다음 24시간 실제 기록과 비교")
add_text(doc, "후보 관계는 예측 시점 직전 4시간에 나타난 방향성 사용자쌍이다. 정답은 다음 24시간에 A에서 B로 향하는 사건이 하나라도 다시 관측되면 1, 관측되지 않으면 0이다. 행동 유형은 같을 필요가 없다. 예를 들어 과거 c2a이고 미래 a2q여도 방향이 A에서 B로 같으면 양성이다. A에서 B와 B에서 A는 서로 다른 관계로 취급한다.")

add_heading(doc, "1.2 간단한 예시", 2)
add_text(doc, "다섯 명의 사용자가 있다고 가정하자. 예측 직전 4시간에 A가 B에게 두 번, C가 B에게 한 번, D가 E에게 한 번 상호작용했다. 그러면 A에서 B, C에서 B, D에서 E가 예측 후보가 된다. 다음 24시간에 A에서 B로 향하는 사건이 다시 관측되고 나머지 두 관계가 관측되지 않았다면 A에서 B는 양성이고 나머지는 관측 음성이다. 모델은 각 후보에 순위 점수를 부여한다. 보정된 분류 모델의 출력만 확률로 해석한다.")

add_heading(doc, "1.3 연구 질문", 2)
add_bullets(doc, [
    "과거 상호작용 기록으로 다음 24시간 재연결을 무작위보다 잘 예측할 수 있는가",
    "최근성 빈도 행동 유형 네트워크 구조 중 어떤 정보가 도움이 되는가",
    "제한된 구현 조건에서 GraphSAGE와 시간 신경망 파이프라인을 실행할 수 있는가",
    "같은 분석 절차를 AskUbuntu에 별도로 적용했을 때 비슷한 방향이 나타나는가",
])

add_heading(doc, "1.4 프로젝트 범위", 2)
add_text(doc, "수행 기간이 2개월인 학부 팀프로젝트라는 점을 고려해, 새로운 거대 모델을 개발하기보다 데이터 처리부터 시간 분할, 기준선 비교, GNN 실험, 오차 분석까지 한 흐름을 완성하는 데 초점을 두었다. 결과가 좋지 않은 모델도 제외하지 않고 왜 실패했는지 분석했다.")


# 2장
add_heading(doc, "2 프로젝트가 발전한 과정", 1)
add_heading(doc, "2.1 인터넷 밈과 X 데이터 단계", 2)
add_text(doc, "처음에는 X에서 럭키비키와 같은 인터넷 밈이 어떤 사용자 관계를 따라 퍼지는지 분석하려 했다. 한 사용자가 밈을 퍼뜨렸을 때 다른 밈도 퍼뜨릴 가능성이 높은지, 연결 중심성과 매개 중심성이 확산과 관련되는지 살펴보는 방향이었다. 이를 위해 X 검색 결과와 리포스트 관계를 수집하는 앱을 만들었다.")
add_text(doc, "하지만 X API는 추가 페이지를 수집할 때 비용이 계속 발생했고 HTTP 402 응답도 나타났다. 수집된 계정 수와 실제 리포스트 관계에 사용할 수 있는 계정 수가 달랐으며, 같은 데이터를 팀원이 다시 얻기 어렵다는 문제도 있었다. 기술 오류는 수정할 수 있었지만 비용과 재현성 문제 때문에 공개 데이터로 방향을 바꾸었다.")

add_heading(doc, "2.2 Higgs Twitter 단계", 2)
add_text(doc, "다음에는 공개된 Higgs Twitter 리트윗 사건을 시간순으로 나누었다. 일부 시점까지만 모델에 보여주고 그 이후에 실제로 발생한 관계를 정답으로 사용했다. 이 단계에서 미래 데이터를 미리 보는 시간 누수를 막는 기본 구조가 만들어졌다.")
add_text(doc, "동료 실험은 전체 리트윗 네트워크를 정적 그래프로 만든 뒤 중심성과 그래프 오토인코더를 이용해 숨긴 링크를 복원했다. 기존 시간 실험은 직전 관계가 다음 24시간에 다시 나타나는지를 HGB와 TGN으로 예측했다. 두 실험은 정답과 난도가 달라 점수를 바로 비교할 수 없었기 때문에, 같은 시간 분할과 같은 후보 관계를 사용하는 방식으로 통합했다.")
add_text(doc, "통합 Higgs 실험에서는 균형 표본에서 History HGB의 AP가 0.7597로 가장 높았지만 자연 후보 전체에서는 PairFrequency AP가 0.0776, History HGB가 0.0716이었다. 균형 표본의 높은 점수가 실제 희소한 분포의 성능을 과장할 수 있다는 점을 확인했다. 또한 Higgs가 약 7일 분량이라 독립적인 24시간 평가 구간과 양성 사례를 충분히 확보하기 어려웠다.")

add_heading(doc, "2.3 Stack Exchange 데이터로 전환", 2)
add_text(doc, "더 긴 기간의 실제 사건이 필요해 MathOverflow와 AskUbuntu로 데이터셋을 바꾸었다. 이 데이터에는 답변과 댓글 관계가 시간과 함께 기록돼 있어 과거와 미래를 명확히 나눌 수 있다. MathOverflow의 초기 결과는 후보 기간과 특성 선택 과정에서 이미 확인했으므로 탐색적·회고적 결과로 분류한다. 이후 절차를 AskUbuntu에 적용했지만 모델은 AskUbuntu 자료로 다시 학습했으므로 완전한 독립 검증이나 제로샷 전이는 아니다. Super User는 원자료를 받기 전에 별도 24시간 후보 프로토콜과 성공 기준을 고정했다.")

add_table(
    doc,
    ["단계", "진행한 일", "발견한 문제", "다음 결정"],
    [
        ["인터넷 밈", "X 리포스트 확산과 중심성 분석", "밈 정의와 확산 경로 증거가 불명확", "특정 관계의 시간 예측으로 질문 구체화"],
        ["X 수집 앱", "검색 페이지와 사용자 관계 수집", "비용 재현성 사용 가능한 관계 부족", "공개 데이터 사용"],
        ["Higgs Twitter", "과거 그래프로 다음 24시간 재연결 예측", "기간이 짧고 양성이 적음", "장기 시간 네트워크 탐색"],
        ["초기 실험 통합", "동료의 정적 GAE와 시간 모델을 같은 조건에서 비교", "균형 평가와 실제 분포의 차이", "자연 후보 AP를 주 지표로 채택"],
        ["최종 실험", "MathOverflow 탐색과 AskUbuntu 반복 및 Super User 별도 확인", "완전 미접촉 4시간 테스트 부재", "주장 범위와 한계 공개"],
    ],
    widths=[2.7, 4.6, 4.4, 4.5],
    font_size=8.5,
)

doc.add_page_break()


# 3장
add_heading(doc, "3 데이터와 정답 만들기", 1)
add_heading(doc, "3.1 데이터셋", 2)
add_text(doc, "세 데이터셋은 Stack Exchange 계열 사이트에서 발생한 사용자 간 사건을 시간순으로 기록한다. 한 행은 누가 누구에게 어떤 유형의 상호작용을 언제 했는지를 나타낸다. 아래 수치는 자기 자신에게 향한 사건을 제거한 뒤 실제 특성 생성에 사용한 사건 수이다.")
add_table(
    doc,
    ["데이터", "전체 사건", "a2q", "c2a", "c2q", "역할"],
    [
        ["MathOverflow", "390441", "104138", "139654", "146649", "주 실험"],
        ["AskUbuntu", "726661", "257305", "253861", "215495", "재현 실험"],
        ["Super User", "1108739", "399420", "388751", "320568", "추가 확인"],
    ],
    widths=[3.2, 2.7, 2.4, 2.4, 2.4, 3.1],
    font_size=8.8,
)
add_text(doc, "a2q는 답변자에서 질문자로 향하는 관계이다. c2a는 댓글 작성자에서 답변자로, c2q는 댓글 작성자에서 질문자로 향하는 관계이다. 세 유형을 하나의 방향성 사건 목록으로 합치되, 특성을 만들 때는 유형별 횟수도 따로 보존했다.")
add_text(doc, "공개 SNAP 파일에는 사용자 두 명, 발생 시각, 행동 유형만 있고 게시물·질문·답변 식별자는 없다. 따라서 같은 게시물에서 이어진 대화와 다른 게시물에서 다시 만난 상호작용을 분리할 수 없다. 이 보고서의 ‘재발생’은 지속적인 사회적 관계가 아니라 데이터에 기록된 방향성 사건의 재관측을 뜻한다.")

add_heading(doc, "3.2 후보와 정답", 2)
add_bullets(doc, [
    "MathOverflow와 AskUbuntu 주 실험의 후보는 예측 직전 4시간에 나타난 방향성 사용자쌍이다.",
    "Super User 별도 확인 실험은 사전 고정된 직전 24시간 후보를 사용하므로 주 실험과 AP를 직접 비교하지 않는다.",
    "양성은 행동 유형과 관계없이 같은 방향의 사용자쌍이 다음 24시간에 다시 관측된 경우이다.",
    "음성은 다음 24시간에 재발생이 관측되지 않은 후보이며 절대적인 관계 부재를 뜻하지 않는다.",
    "모든 특성은 예측 시점 이전 기록만 사용해 계산한다.",
])

add_heading(doc, "3.3 시간 분할", 2)
add_text(doc, "무작위로 행을 섞으면 미래 사건이 학습 데이터에 들어갈 수 있다. 이를 막기 위해 학습, 튜닝, 확률 보정, 테스트 구간을 시간순으로 배치했다. 각 앞 구간의 마지막 정답 창이 다음 구간 시작 전 끝나도록 cutoff+24시간 이하 다음 구간 시작 규칙을 적용했다. 테스트 폴드끼리도 겹치지 않는다.")
add_text(doc, "데이터 마지막 시점에서는 다음 24시간을 관찰할 수 없으므로 정답 구간이 완전히 존재하는 cutoff만 사용했다. 경계 시각의 사건은 cutoff 이전이면 특성에, cutoff 이상 cutoff+24시간 미만이면 정답에 포함한다. 같은 사건이 입력과 정답에 동시에 사용되지 않는다.")
add_table(
    doc,
    ["데이터", "폴드", "테스트 기간 UTC", "후보", "양성"],
    [
        ["MathOverflow", "0", "2013-01-02~2013-08-22", "7317", "830"],
        ["MathOverflow", "1", "2013-08-22~2014-04-11", "6227", "783"],
        ["MathOverflow", "2", "2014-04-11~2014-11-29", "6136", "726"],
        ["MathOverflow", "3", "2014-11-29~2015-07-18", "5622", "701"],
        ["MathOverflow", "4", "2015-07-18~2016-03-06", "6003", "747"],
        ["AskUbuntu", "0", "2013-05-17~2013-12-08", "15110", "868"],
        ["AskUbuntu", "1", "2013-12-08~2014-07-01", "17394", "1181"],
        ["AskUbuntu", "2", "2014-07-01~2015-01-21", "15252", "1135"],
        ["AskUbuntu", "3", "2015-01-21~2015-08-14", "18566", "1458"],
        ["AskUbuntu", "4", "2015-08-14~2016-03-06", "17989", "1445"],
    ],
    widths=[3.2, 1.4, 6.2, 2.6, 2.6],
    font_size=8.0,
)

add_heading(doc, "3.4 분석 단계와 누수 감사", 2)
add_text(doc, "직전 4시간 후보는 초기 탐색을 거쳐 정한 연구 범위이다. 따라서 MathOverflow 결과를 사전등록 확인 실험으로 표현하지 않고 회고적 시간 평가로 분류한다. AskUbuntu는 고정한 절차를 해당 데이터에서 다시 학습한 재현 실험이며, MathOverflow 모델을 그대로 옮긴 제로샷 전이 실험은 아니다.")
add_text(doc, "감사 프로그램은 데이터셋마다 원자료 정답 1000건을 재계산하고 255개 모델 실행의 후보, 정답, 저장 예측과 지표를 다시 확인했다. 두 데이터셋 모두 과거 후보 존재, 완전한 24시간 정답 창, 분할 경계의 24시간 purge, 비중첩 테스트 기간, 모델 간 동일 후보, 저장 지표 재계산 검사를 통과했다. 이 결과는 구현된 검사 범위에서 누수가 발견되지 않았다는 뜻이며 모든 가능한 오류가 없음을 보증하지는 않는다.")


# 4장
add_heading(doc, "4 모델이 보는 정보", 1)
add_heading(doc, "4.1 입력 특성", 2)
add_text(doc, "모델은 사용자 이름이나 미래 게시물을 보는 것이 아니라 예측 시점 이전에 계산한 숫자 특성을 사용한다. Full LightGBM에는 총 78개의 시간, 행동 유형, 네트워크 특성이 들어간다. 공통 이웃과 Jaccard도 전체 기간 그래프가 아니라 각 cutoff 직전 30일 사건으로 다시 계산한다.")
add_table(
    doc,
    ["특성 묶음", "쉬운 설명", "예시"],
    [
        ["관계 이력", "두 사용자가 과거에 얼마나 자주 연결됐는지", "전체 횟수 최근 4시간 횟수"],
        ["최근성", "마지막 연결 이후 시간이 얼마나 지났는지", "마지막 사건 후 경과 시간"],
        ["행동 유형", "답변과 댓글 중 어떤 방식으로 연결됐는지", "a2q c2a c2q별 횟수"],
        ["사용자 활동", "각 사용자가 최근 얼마나 활발했는지", "보낸 사건 수 받은 사건 수"],
        ["네트워크 구조", "두 사용자 주변에 공통으로 연결된 사람이 있는지", "공통 이웃 Jaccard 중심성"],
        ["상호성", "반대 방향의 관계도 과거에 있었는지", "B에서 A로 향한 최근 사건"],
    ],
    widths=[3.2, 7.3, 5.7],
    font_size=8.8,
)

add_heading(doc, "4.2 비교 모델", 2)
add_table(
    doc,
    ["모델", "역할", "이 프로젝트에서의 의미"],
    [
        ["Pair frequency", "두 사용자 사이의 과거 횟수만 사용", "가장 단순한 기준선"],
        ["History HGB", "빈도와 최근성 등 핵심 이력 사용", "강한 전통적 기준선"],
        ["Full LightGBM", "78개 특성을 함께 학습", "주 모델"],
        ["GraphSAGE", "주변 사용자 구조를 임베딩으로 학습", "정적 GNN 비교"],
        ["TGN recurrence", "시간에 따라 변하는 관계를 신경망으로 학습", "시간 GNN 비교"],
    ],
    widths=[4.0, 6.2, 6.0],
    font_size=8.8,
)
add_text(doc, "모델이 복잡하다고 반드시 더 좋은 것은 아니다. 단순 기준선과 복잡한 모델을 같은 후보, 같은 시간 분할, 같은 정답으로 비교해야 추가 복잡성이 실제로 도움이 됐는지 판단할 수 있다.")
add_text(doc, "모든 모델은 같은 후보와 시간 분할, 정답으로 평가했다. 그러나 입력 정보와 최적화 예산은 같지 않다. Full LightGBM은 78개 수작업 특성과 폴드별 Optuna 16회 탐색을 사용한 반면, GraphSAGE와 TGN은 사건 흐름과 최근 이웃 구조를 직접 입력받고 최대 4 epoch에서 선택했다. 따라서 이 절의 목적은 GNN의 성능 우열을 검정하는 것이 아니라 GPU 기반 그래프 학습 파이프라인의 구현 가능성과 현재 설정의 결과를 기록하는 데 있다.")
add_table(
    doc,
    ["모델", "주 입력", "튜닝과 학습", "해석 범위"],
    [
        ["Full LightGBM", "78개 과거 기반 표형 특성", "폴드별 Optuna 16회", "주 기준 모델"],
        ["GraphSAGE", "최근 이웃 10개와 노드 활동 수", "차원 32, 최대 4 epoch", "제한된 정적 GNN 구현"],
        ["TGN recurrence", "사건 유형 방향 시간과 메모리", "차원 32, 최대 4 epoch", "제한된 시간 GNN 구현"],
    ],
    widths=[3.4, 5.1, 3.7, 4.0],
    font_size=8.5,
)


# 5장
add_heading(doc, "5 평가 지표를 읽는 방법", 1)
add_heading(doc, "5.1 양성과 음성", 2)
add_text(doc, "양성은 다음 24시간에 같은 방향의 사건이 관측된 후보이고 음성은 관측되지 않은 후보이다. MathOverflow 4시간 후보의 양성률은 약 12.1퍼센트이고 AskUbuntu는 약 7.2퍼센트이다. Super User의 약 3.7퍼센트는 24시간 후보 전체의 값이므로 직접 비교하지 않는다. 관측되지 않은 사건과 삭제·누락된 사건을 구분할 수 없으므로 음성은 ‘관측 음성’으로 해석한다.")

add_heading(doc, "5.2 AP와 AUC와 F1", 2)
add_table(
    doc,
    ["지표", "무엇을 보는가", "이번 프로젝트에서의 사용"],
    [
        ["AP", "높은 점수를 받은 후보에 실제 양성이 얼마나 잘 모이는지", "주 지표"],
        ["AUC", "임의의 양성이 임의의 음성보다 높은 점수를 받을 가능성", "보조 지표"],
        ["F1", "하나의 임계값에서 정밀도와 재현율의 균형", "보조 지표"],
        ["P@5", "각 예측 시점에서 점수가 높은 5개 중 실제 양성 비율", "운영 관점의 상위 후보 평가"],
        ["R@5", "전체 양성 중 상위 5개 후보로 찾은 비율", "상위 후보의 포착 범위"],
        ["양성률", "전체 후보 중 실제 재연결 비율", "무작위 AP의 기준"],
    ],
    widths=[2.6, 7.8, 5.8],
    font_size=9.0,
)
add_text(doc, "AP는 불균형 자료에서 모델의 순위 성능을 보기 좋다. 무작위 점수의 기대 AP는 대략 양성률과 같다. AP를 양성률로 나눈 1.87배와 3.02배는 이해를 돕는 서술적 비율이며 표준화된 효과 크기는 아니다.")
add_text(doc, "P@5는 각 일별 cutoff에서 상위 5개를 선택한 뒤 동점에 기대값을 주고, 전체 cutoff의 적중 수를 전체 선택 수로 나눈 micro 집계이다. R@5는 같은 적중 수를 모든 cutoff의 전체 양성 수로 나눈다. 후보가 5개보다 적으면 가능한 후보만 사용한다. 무작위 기대 적중 수 0.60개와 0.36개는 전체 테스트 양성률에 5를 곱한 단순 참고값이다. cutoff마다 후보 수와 양성률이 다르므로 정확한 cutoff별 무작위화 검정으로 해석하지 않는다.")
add_text(doc, "F1 임계값과 로지스틱 확률 보정은 테스트보다 앞선 보정 구간에서 폴드별로 선택했다. AP와 AUC는 순위 지표이므로 보정 여부와 관계없이 해석할 수 있지만, ‘확률’이라는 표현은 보정된 분류 모델 출력에만 사용한다. AP, P@5, R@5가 높아도 실제 서비스 비용과 편익은 별도 검증이 필요하다.")
add_text(doc, "초기 Higgs 실험의 균형 평가에서는 F1이 약 0.667로 높게 나왔다. 그러나 양성과 음성을 같은 수로 뽑으면 실제보다 문제가 쉬워진다. 최종 보고서는 자연 후보 전체의 AP를 중심으로 보고하고 F1은 보조적으로만 해석한다. F1이 0.5보다 낮다고 해서 실험이 틀렸다는 뜻은 아니다.")

# 6장
add_heading(doc, "6 탐색 및 반복 실험 결과", 1)
add_heading(doc, "6.1 MathOverflow 탐색 실험", 2)
add_table(
    doc,
    ["모델", "AP", "AUC", "F1", "해석"],
    [
        ["Full LightGBM", "0.2260", "0.6995", "0.3008", "주 모델"],
        ["TGN recurrence", "0.2103", "0.6797", "0.2872", "주 모델보다 낮음"],
        ["History HGB", "0.2012", "0.6493", "0.2707", "강한 기준선"],
        ["GraphSAGE", "0.1819", "0.6323", "0.2592", "구조만으로는 부족"],
        ["Pair frequency", "0.1528", "0.5791", "0.2429", "단순 기준선"],
        ["양성률", "0.1210", "해당 없음", "해당 없음", "무작위 AP 기준"],
    ],
    widths=[3.9, 2.0, 2.0, 2.0, 6.3],
    font_size=8.5,
)
add_text(doc, "표의 AP, AUC, F1은 각 폴드에서 시드 예측을 평균한 뒤 계산한 지표를 다시 다섯 시간 폴드에 걸쳐 평균한 값이다. pooled 지표가 아니며 폴드별 AP는 6.4절에 별도로 제시한다.")
add_text(doc, "Full LightGBM의 AP는 0.2260으로 양성률 0.1210보다 약 1.87배 높았다. F1은 0.3008, 정밀도는 0.2120, 재현율은 0.5250이었다. 실제 재연결의 절반 정도를 찾았지만 양성으로 판단한 후보 중에는 재연결하지 않은 관계도 많았다.")
add_figure(doc, PRIMARY_MO_FIGURE, "1 MathOverflow 주 모델별 AP 비교")

doc.add_page_break()
add_heading(doc, "6.2 AskUbuntu 도메인 재학습 반복 실험", 2)
add_table(
    doc,
    ["모델", "AP", "AUC", "F1", "해석"],
    [
        ["Full LightGBM", "0.2183", "0.7856", "0.2853", "주 모델"],
        ["History HGB", "0.1939", "0.7423", "0.2645", "강한 기준선"],
        ["TGN recurrence", "0.1900", "0.7633", "0.2574", "주 모델보다 낮음"],
        ["GraphSAGE", "0.1636", "0.7202", "0.2291", "주 모델보다 낮음"],
        ["양성률", "0.0722", "해당 없음", "해당 없음", "무작위 AP 기준"],
    ],
    widths=[3.9, 2.0, 2.0, 2.0, 6.3],
    font_size=8.5,
)
add_text(doc, "표의 AP, AUC, F1은 MathOverflow 표와 같은 방식의 다섯 시간 폴드 평균이다. 한 번에 모든 테스트 예측을 합친 pooled 결과와 구분한다.")
add_text(doc, "AskUbuntu에서 Full LightGBM의 AP는 0.2183으로 양성률 0.0722의 약 3.02배였다. MathOverflow에서 정한 절차를 적용하되 AskUbuntu 자료에서 모델을 새로 학습했을 때 같은 방향의 결과가 관찰됐다. 이는 동일 계열 플랫폼에서 수행한 도메인별 재학습 반복이며, MathOverflow 모델을 그대로 적용한 독립 외부 검증이나 플랫폼 간 전이 검증은 아니다.")
add_figure(doc, PRIMARY_AU_FIGURE, "2 AskUbuntu 주 모델별 AP 비교")

add_heading(doc, "6.3 상위 5개 후보 평가", 2)
add_table(
    doc,
    ["데이터", "양성률", "P@5", "R@5", "해석"],
    [
        ["MathOverflow", "0.1210", "0.2292", "0.3521", "상위 5개 중 평균 약 1.15개 적중"],
        ["AskUbuntu", "0.0722", "0.2608", "0.2223", "상위 5개 중 평균 약 1.30개 적중"],
    ],
    widths=[3.2, 2.3, 2.3, 2.3, 6.1],
    font_size=8.9,
)
add_text(doc, "전체 양성률을 사용한 단순 참고 계산에서는 MathOverflow의 무작위 기대 적중 수가 약 0.60개이고 모델의 상위 5개에서는 약 1.15개였다. AskUbuntu에서는 참고값 약 0.36개와 모델 결과 약 1.30개였다. cutoff별 후보 수와 양성률 차이를 반영한 정확한 무작위화 검정은 아니므로, 이 비교는 상위 후보에 신호가 집중됐는지를 설명하는 보조 근거로만 사용한다. 또한 모델이 선택한 후보의 다수는 재발생하지 않았다.")

add_heading(doc, "6.4 어떤 정보가 도움이 됐는가", 2)
add_text(doc, "5개 폴드 평균에서는 Full LightGBM이 History HGB보다 높았고 행동 유형 특성을 제거하면 성능이 낮아졌다. 이 차이의 불확실성은 시드 점수를 먼저 평균한 뒤 7일 시간 블록을 1000회 재표집해 평가했다. 아래 값은 모든 비중첩 테스트 예측을 합친 pooled AP 차이이므로 앞 표의 폴드 평균 AP를 단순히 뺀 값과 약간 다르다.")
add_table(
    doc,
    ["데이터", "시간 폴드", "Full LightGBM AP", "History HGB AP", "차이"],
    [
        ["MathOverflow", "0", "0.2088", "0.1886", "+0.0201"],
        ["MathOverflow", "1", "0.2218", "0.1912", "+0.0306"],
        ["MathOverflow", "2", "0.2510", "0.2150", "+0.0360"],
        ["MathOverflow", "3", "0.2285", "0.2069", "+0.0216"],
        ["MathOverflow", "4", "0.2199", "0.2043", "+0.0156"],
        ["AskUbuntu", "0", "0.1856", "0.1658", "+0.0199"],
        ["AskUbuntu", "1", "0.2072", "0.1818", "+0.0254"],
        ["AskUbuntu", "2", "0.2369", "0.2021", "+0.0349"],
        ["AskUbuntu", "3", "0.2109", "0.1914", "+0.0194"],
        ["AskUbuntu", "4", "0.2510", "0.2282", "+0.0228"],
    ],
    widths=[3.2, 2.3, 3.8, 3.5, 2.5],
    font_size=8.1,
)
add_text(doc, "Full LightGBM과 History HGB의 AP 차이는 두 데이터셋의 다섯 시간 폴드에서 모두 양의 방향이었다. 이 표는 폴드별 변동을 숨기지 않기 위한 기술적 요약이며, 탐색 과정에서 선택된 전체 분석 절차의 사후 선택 문제를 제거하지는 않는다.")
add_table(
    doc,
    ["데이터", "비교", "AP 차이", "95퍼센트 시간 블록 구간"],
    [
        ["MathOverflow", "Full LightGBM - History HGB", "+0.0258", "+0.0214~+0.0305"],
        ["MathOverflow", "Full LightGBM - 유형 제거", "+0.0217", "+0.0172~+0.0265"],
        ["AskUbuntu", "Full LightGBM - History HGB", "+0.0238", "+0.0197~+0.0280"],
        ["AskUbuntu", "Full LightGBM - 유형 제거", "+0.0187", "+0.0155~+0.0220"],
    ],
    widths=[3.1, 6.3, 2.6, 4.2],
    font_size=8.4,
)
add_text(doc, "History HGB와 유형 제거 모델에 대한 차이는 선택한 7일 블록 가정에서 0보다 높았다. 시간감쇠 특성을 제거한 모델의 평균 AP는 오히려 Full LightGBM보다 조금 높았으므로 최근성이 독립적으로 가장 중요한 원인이라고 결론 내리지 않는다. 안전한 해석은 선택한 후보와 특성 구성에서 관계 이력과 행동 유형을 포함한 특성 조합이 추가 순위 신호를 보였다는 것이다.")

add_heading(doc, "6.5 시간 블록 길이 민감도", 2)
add_text(doc, "주 비교인 Full LightGBM과 History HGB의 pooled AP 차이에 대해 이미 계산된 7일, 14일, 28일 블록 결과를 함께 제시한다. 모든 계산은 모델 시드 점수를 먼저 평균한 뒤 시간 블록을 1000회 재표집했다. 따라서 표의 구간은 시간 의존성에 대한 민감도는 보여주지만 학습 시드 변동을 동시에 재표집한 구간은 아니다.")
add_table(
    doc,
    ["데이터", "블록 길이", "pooled AP 차이", "95퍼센트 구간"],
    [
        ["MathOverflow", "7일", "+0.0258", "+0.0214~+0.0305"],
        ["MathOverflow", "14일", "+0.0258", "+0.0213~+0.0301"],
        ["MathOverflow", "28일", "+0.0258", "+0.0207~+0.0310"],
        ["AskUbuntu", "7일", "+0.0238", "+0.0197~+0.0280"],
        ["AskUbuntu", "14일", "+0.0238", "+0.0197~+0.0284"],
        ["AskUbuntu", "28일", "+0.0238", "+0.0195~+0.0281"],
    ],
    widths=[4.0, 3.0, 4.0, 5.2],
    font_size=8.7,
)
add_text(doc, "세 블록 길이에서 구간이 0보다 높았다는 사실은 선택한 7일 값 하나에만 결과가 의존하지 않았음을 보여준다. 다만 1일 블록과 시드·시간을 함께 반영한 계층적 재표집은 수행하지 않았고, MathOverflow 자체가 탐색 자료였으므로 이를 보편적 우월성의 확증으로 해석하지 않는다.")

# 7장
add_heading(doc, "7 추가 실험과 시행착오", 1)
add_heading(doc, "7.1 후보 기간을 늘린 실험", 2)
add_text(doc, "직전 4시간 후보는 초기 탐색 과정에서 정한 조건이며 MathOverflow 결과를 보기 전에 외부에 사전등록한 값은 아니다. 따라서 4시간을 보편적으로 최적이라고 주장하지 않는다. 민감도 분석에서는 후보 기간을 12시간과 24시간으로 늘렸다. MathOverflow의 Full LightGBM 평균 AP는 0.2260, 0.1950, 0.1682였고 AskUbuntu도 0.2183, 0.1795, 0.1533이었다.")
add_text(doc, "후보 기간을 늘리면 오래전에 한 번 연결되고 다시 나타날 가능성이 낮은 관계가 추가되고 양성률도 달라진다. 따라서 서로 다른 후보 모집단의 AP만으로 4시간 모델이 더 우수하다고 결론 내릴 수 없다. 동일한 기존 4시간 후보만 다시 평가하면 AP가 거의 유지됐다. 이 보고서는 4시간을 단기 재발생의 운영적 정의로 사용하고 12시간과 24시간 결과를 민감도 분석으로 함께 공개한다.")
add_figure(doc, RESULTS / "lookback_density_v1" / "mathoverflow" / "lookback_ap.png", "3 후보 기간과 학습 간격에 따른 AP")

add_heading(doc, "7.2 어떤 관계를 놓치고 잘못 예측했는가", 2)
add_text(doc, "오차 분석은 주 실험과 같은 직전 4시간 후보와 24시간 학습 간격에서 시드 2026의 다섯 테스트 폴드를 합쳐 수행했다. 독립적인 가설 검정이 아니라 어떤 사례가 어려웠는지 설명하는 탐색적 분석이다.")
add_text(doc, "마지막 상호작용이 0시간에서 1시간 전인 관계의 놓침률은 약 27퍼센트였지만 1시간에서 4시간 전인 관계는 약 59퍼센트였다. 누적 상호작용이 한 번뿐인 관계와 a2q 유형도 많이 놓쳤고, 예측 전에 반대 방향 상호작용이 있었던 관계는 상대적으로 덜 놓쳤다.")
add_figure(doc, ASSETS / "error_miss_rate_ko.png", "4 MathOverflow 주 설정의 놓침률")
add_text(doc, "거짓발견률은 실제 양성 대신 예측 양성을 분모로 사용하므로 놓침률과 직접 비교하지 않았다. 모든 하위 집단에서 거짓발견률이 높았으며 특히 a2q에서 높았다. 이는 현재 특성 조합과 선택된 임계값에서 정확한 상대와 행동 유형을 구분하는 일이 어려웠음을 보여준다. 특정 특성이 오류를 발생시켰다는 인과적 의미는 아니다.")
add_figure(doc, ASSETS / "error_false_discovery_ko.png", "5 MathOverflow 주 설정의 거짓발견률")

add_heading(doc, "7.3 Super User 사전 고정 추가 확인", 2)
add_text(doc, "기존 두 데이터셋 결과를 확인한 뒤 같은 테스트 자료에서 모델을 계속 바꾸는 문제를 줄이기 위해 Super User 원자료를 내려받기 전에 후보 범위, 시간 폴드, 성공 기준과 모델을 파일로 고정했다. 이 실험은 직전 24시간에 관측된 후보 전체를 사용하고 학습 cutoff는 6시간 간격이다. 따라서 직전 4시간 후보를 사용한 주 실험과 AP를 직접 비교할 수 없다.")
add_text(doc, "같은 78개 특성을 사용한 Baseline Full LightGBM의 AP는 0.1393으로 24시간 후보 양성률 0.0366보다 높았다. 재연결 예측을 활동 여부와 상대 선택으로 나눈 최근성·유형 전문가 모델의 pooled AP 증가는 0.00027이었고 95퍼센트 7일 블록 구간은 -0.00089에서 0.00151이었다. 사전 고정 성공 조건을 충족하지 못했으므로 새 모델의 우월성을 확인하지 못했다. 이는 다른 사이트에서의 도메인별 학습 결과이며 제로샷 전이 검증은 아니다.")
add_figure(doc, ASSETS / "superuser_confirmation_ap_ko.png", "6 Super User 추가 확인 모델 비교")

add_heading(doc, "7.4 제한된 구현 조건의 GNN과 GPU 랭킹 탐색", 2)
add_text(doc, "GraphSAGE와 TGN을 추가하고 GPU 학습 경로도 만들었다. 두 모델은 CUDA에서 실행됐지만 최대 4 epoch로 제한됐고 LightGBM과 동일한 78개 특성을 입력받지도 않았다. 현재 구현에서는 Full LightGBM보다 낮은 수치를 보였으나 이 결과는 GNN의 성능 검정이 아니라 그래프 학습 파이프라인의 구현 가능성을 확인한 탐색 결과다. 이후 상대 선택을 직접 학습하는 GPU RankNet도 탐색했지만 Super User AP는 0.1285로 Baseline Full LightGBM 0.1393보다 낮았다.")

add_heading(doc, "7.5 결합 모델을 핵심 결과에서 제외한 이유", 2)
add_text(doc, "개발 과정에서는 LightGBM의 predict_proba 출력과 TGN logit에 sigmoid를 적용한 출력을 선형 결합하는 Hybrid도 시험했다. 두 출력은 모두 0과 1 사이였지만 결합 전에 동일한 기준으로 공동 확률 보정하거나 분포를 정규화하지 않았다. 따라서 선택된 결합 가중치를 두 모델의 확률적 기여도나 직접적인 중요도로 해석할 수 없다.")
add_text(doc, "가중치는 각 폴드와 시드의 튜닝 구간에서만 선택했고 결합 뒤에는 별도의 확률 보정을 적용했으므로 테스트 정답을 직접 사용한 결합은 아니다. 그러나 AP 개선 폭이 MathOverflow 약 0.0016, AskUbuntu 약 0.0012로 작고 F1도 Full LightGBM보다 낮았다. 이에 따라 Hybrid는 구현 과정의 탐색 기록으로만 공개하고 주 결과표, 주 그림, 모델 우열 비교와 최종 결론에서는 제외했다.")


# 8장
add_heading(doc, "8 결과 해석", 1)
add_heading(doc, "8.1 확인된 점", 2)
add_bullets(doc, [
    "직전 4시간에 이미 관측된 방향성 사용자쌍 안에서는 과거 정보가 다음 24시간 관측 재발생의 순위 신호로 작동했다.",
    "상위 5개 후보 평가는 두 데이터셋에서 전체 양성률 기반 참고값보다 높았지만 정확한 cutoff별 무작위화 검정은 아니다.",
    "관계 이력과 행동 유형을 포함한 특성 조합은 선택한 설정에서 History HGB와 유형 제거 모델보다 높은 AP를 보였다.",
    "같은 절차를 AskUbuntu에서 별도로 재학습했을 때도 비슷한 방향이 관찰됐지만 독립 외부 검증은 아니다.",
    "Super User의 별도 24시간 후보에서도 Baseline Full LightGBM AP가 해당 후보 양성률보다 높았다.",
])

add_heading(doc, "8.2 확인되지 않은 점", 2)
add_bullets(doc, [
    "동일한 입력과 튜닝 예산을 사용한 공정한 GNN 대 LightGBM 우열은 검증하지 않았다.",
    "활동 여부와 상대 선택을 나눈 전문가 모델의 개선은 Super User에서 확인되지 않았다.",
    "GPU RankNet도 강한 기준선을 넘지 못했다.",
    "신규 링크 예측과 MathOverflow 모델의 다른 사이트 제로샷 전이 성능은 평가하지 않았다.",
    "모든 분석 선택과 무관하게 남겨 둔 직전 4시간의 완전 미접촉 최종 테스트는 확보하지 못했다.",
])

add_heading(doc, "8.3 결과의 의미", 2)
add_text(doc, "이 프로젝트의 목적은 높은 F1 하나를 만드는 것이 아니라 시간 순서를 지키면서 자연 후보 전체의 순위 신호를 평가하는 것이다. AP와 P@5가 양성률 기반 참고값보다 높고 유형 제거 실험에서도 선택한 시간 블록 가정 아래 양의 차이가 관찰됐다. 따라서 본 자료와 분석 설정 안에서는 과거 사용자 상호작용에 조건부 단기 재발생 후보를 순위화하는 정보가 있었다고 해석한다. 이는 탐색적 예측 연관이며 행동 유형이나 최근성이 재발생을 발생시킨다는 인과적 증거가 아니다.")
add_text(doc, "오류 분석에서는 정확한 상대를 고르는 문제가 크게 남았다. 질문 태그, 주제 유사도, 게시물 텍스트 같은 정보는 향후 검증할 수 있는 추가 가설이지만 현재 결과만으로 성능 향상을 보장할 수 없다. 실제 서비스에서 필요한 비용과 편익 기준도 정하지 않았으므로 배포 가능한 수준의 유용성을 입증한 것으로 해석하지 않는다.")

add_heading(doc, "8.4 한계", 2)
add_bullets(doc, [
    "세 데이터셋은 모두 Stack Exchange 계열이므로 다른 SNS에 바로 일반화할 수 없다.",
    "AskUbuntu와 Super User에서는 해당 데이터로 다시 학습했으며 MathOverflow 모델을 그대로 옮긴 전이 실험은 아니다.",
    "후보를 최근 관계로 제한했기 때문에 한 번도 연결되지 않은 완전한 신규 링크는 예측하지 않는다.",
    "원자료에 게시물 ID가 없어 같은 스레드의 연속 활동과 다른 게시물에서의 재접촉을 구분하지 못한다.",
    "재발생이 기록되지 않은 후보를 관측 음성으로 처리하므로 삭제·누락 사건과 실제 부재를 구분할 수 없다.",
    "사용자 관계와 시간 정보가 중심이며 게시물의 텍스트 의미는 포함하지 않았다.",
    "LightGBM과 GNN의 후보와 정답은 같지만 튜닝과 학습 예산은 동일하지 않아 모델 계열의 일반적 우열을 결론 내릴 수 없다.",
    "P@5는 제한된 후보 확인 상황을 가정한 지표이며 실제 서비스의 비용과 개입 효과는 검증하지 않았다.",
    "4시간 후보는 초기 탐색 뒤 정한 운영적 범위이므로 보편적인 최적값이나 사전등록 효과로 해석하지 않는다.",
    "MathOverflow는 모델과 특성 선택에 사용된 탐색 자료이며 AskUbuntu도 도메인별 재학습을 수행했으므로 완전히 독립적인 확인 자료가 아니다.",
    "7일 14일 28일 블록 결과는 시드 평균 예측을 재표집했기 때문에 학습 시드 변동까지 포함한 신뢰구간이 아니다.",
    "오차 분석은 한 시드의 통합 결과를 설명한 탐색적 분석이므로 인과적 해석을 하지 않는다.",
    "Super User의 개선 모델과 RankNet은 사후 아이디어의 탐색 결과이므로 기본 모델보다 강한 결론으로 해석하지 않는다.",
])


# 9장
add_heading(doc, "9 재현 방법과 프로젝트 구성", 1)
add_heading(doc, "9.1 실행 흐름", 2)
add_bullets(doc, [
    "데이터 파일을 읽고 사용자 관계 사건을 시간순으로 정렬한다.",
    "주 실험의 각 예측 시점에서 직전 4시간의 후보 관계를 만든다.",
    "예측 시점 이전 자료만 사용해 시간 행동 유형 네트워크 특성을 계산한다.",
    "과거 학습·튜닝 구간으로 모델을 정하고 보정 구간에서 확률 보정과 F1 임계값을 선택한다.",
    "뒤쪽 24시간의 실제 사건으로 AP AUC F1을 평가한다.",
    "감사 프로그램으로 원자료 정답, 분할 경계, 저장 예측과 지표를 재검산한다.",
    "모델 차이는 시드를 평균한 뒤 7일 14일 28일 시간 블록 부트스트랩으로 불확실성을 계산한다.",
])

add_heading(doc, "9.2 재현 실행 명령", 2)
add_table(
    doc,
    ["목적", "PowerShell 실행 명령"],
    [
        ["환경 설치", "python -m pip install -r requirements-research.txt"],
        ["전체 주 실험", "python -m research.batch"],
        ["결과 집계", "python -m research.finalize"],
        ["단일 실행 예시", "python -m research.run --dataset mathoverflow --fold 0 --seeds 2026"],
        ["확장 실험", ".\\run_extensions.ps1"],
    ],
    widths=[4.0, 12.2],
    font_size=8.5,
)
add_text(doc, "주 실험은 비중첩 테스트 시간 폴드 5개와 시드 2026, 2027, 2028을 사용한다. 신경망은 CUDA GPU를 사용했고 트리 모델은 CPU를 사용했다. 데이터 URL·SHA-256·시작과 종료 시각, Python과 라이브러리 버전, 78개 특성명, Optuna 탐색 결과, 폴드 경계, 보정 임계값, 모델별 예측값과 실행 시간은 결과 폴더의 manifest와 실행 파일에 저장된다.")

add_heading(doc, "9.3 주요 파일", 2)
add_table(
    doc,
    ["경로", "내용"],
    [
        ["RESEARCH_PROTOCOL_KO.md", "주 실험의 시간 분할 후보 모델 지표 정의"],
        ["research", "주 실험 데이터 처리 특성 모델 평가 코드"],
        ["research_extensions", "후보 기간 민감도와 오차 분석 코드"],
        ["research_confirmation", "Super User 추가 확인 코드"],
        ["research_ranking", "GPU RankNet 탐색 코드"],
        ["results", "표 그림 예측값 요약 파일"],
        ["results/research_v2/*/run_manifest.json", "데이터 해시 환경 버전 78개 특성 목록"],
        ["results/research_v2/*/fold_*/split.json", "학습 튜닝 보정 테스트 경계와 표본 수"],
        ["results/research_v2/*/audit.json", "누수 정답 지표 재검산 결과"],
        ["results/research_v2/*/paired_block_bootstrap.csv", "모델 간 pooled AP 차이와 95퍼센트 구간"],
    ],
    widths=[5.2, 11.0],
    font_size=9.0,
)

add_heading(doc, "9.4 팀원이 결과를 확인하는 순서", 2)
add_text(doc, "처음 보는 팀원은 1장부터 6장까지 읽은 뒤 RESEARCH_PROTOCOL_KO.md에서 세부 조건을 확인하면 된다. 수치 원본은 results 폴더에 있으며 주 실험, 사후 확장 실험, 사전 고정 Super User 실험을 구분해 해석해야 한다. 보고서의 표는 읽기 쉬운 요약이고 독립 재현에는 manifest, split, audit, 예측 파일과 실행 로그를 함께 제공해야 한다.")


# 10장
add_heading(doc, "10 결론", 1)
add_text(doc, "이 팀프로젝트는 직전 4시간에 이미 관측된 방향성 사용자쌍을 대상으로 다음 24시간의 관측된 상호작용 재발생을 순위화했다. 인터넷 밈과 X 수집에서 시작해 Higgs Twitter의 짧은 기간 문제를 확인했고, 더 긴 MathOverflow와 AskUbuntu 자료에서 시간 순서를 지키는 분석 절차를 구축했다. MathOverflow 결과는 후보 기간과 특성을 탐색한 회고적 결과다.")
add_text(doc, "과거 빈도, 최근성, 행동 유형을 포함한 Full LightGBM은 MathOverflow와 AskUbuntu에서 해당 후보 양성률보다 높은 AP와 P@5를 보였다. History HGB 및 유형 제거 모델과의 pooled AP 차이는 시드 평균 예측에 적용한 7일 14일 28일 시간 블록 가정 아래 모두 양의 구간을 보였다. AskUbuntu 결과는 같은 절차를 해당 도메인에서 다시 학습한 반복 결과이며 독립 외부 검증이나 플랫폼 간 전이는 아니다. Super User의 24시간 후보 사전 고정 실험에서는 기본 순위 신호가 관찰됐지만 새 전문가 모델의 개선은 확인되지 않았다.")
add_text(doc, "최종 주 모델은 Full LightGBM으로 정했다. GraphSAGE, TGN과 GPU RankNet은 제한된 구현에서 기준선보다 낮았지만 입력 정보와 튜닝 예산이 달라 GNN의 본질적 열등성으로 해석하지 않는다. 게시물 ID가 없어 같은 스레드의 연속 활동을 분리하지 못하고, 완전히 미접촉인 4시간 최종 테스트도 없다. 따라서 결론은 선택한 자료와 분석 조건에서 최근 상호작용 사용자쌍의 단기 재발생을 순위화하는 예측 신호가 관찰됐다는 범위로 제한한다.")

add_heading(doc, "부록 A Full LightGBM 특성 목록", 1)
add_text(doc, "아래 목록은 주 모델에 실제로 입력된 78개 특성의 정확한 이름이다. 이름의 1h, 4h, 24h, 168h, 720h와 30d는 cutoff 직전의 시간 창을 뜻하고 lifetime은 cutoff 이전 전체 기록을 뜻한다. decay의 시간 값은 최근 사건에 exp(-경과시간/시간상수)를 적용할 때의 시간상수이며 최대 직전 30일 사건을 합산한다. 모든 특성은 예측 cutoff 이전 자료만 사용한다.")
add_table(
    doc,
    ["범주와 개수", "정확한 특성명"],
    [
        ["관계 및 사용자 활동 23개", "pair_count_1h, pair_count_4h, pair_count_12h, pair_count_24h, pair_count_72h, pair_count_168h, pair_count_720h, pair_count_lifetime, reverse_count_4h, reverse_count_24h, reverse_count_168h, reverse_count_720h, reverse_count_lifetime, hours_since_pair_event, pair_age_days, src_out_24h, src_out_168h, src_out_720h, src_in_168h, dst_in_24h, dst_in_168h, dst_in_720h, dst_out_168h"],
        ["30일 네트워크 구조 6개", "src_degree_30d, dst_degree_30d, common_neighbors_30d, jaccard_30d, adamic_adar_30d, preferential_attachment_30d"],
        ["감쇠 및 시간 패턴 19개", "pair_decay_1h, pair_decay_6h, pair_decay_24h, pair_decay_72h, pair_decay_168h, reverse_decay_6h, reverse_decay_24h, reverse_decay_168h, gap_mean_hours, gap_cv, gap_last_hours, burstiness, active_days_30d, count_acceleration, reciprocity_24h, src_acceleration, dst_acceleration, weekday_sin, weekday_cos"],
        ["행동 유형별 특성 27개", "type_a2q_count_4h, type_a2q_count_24h, type_a2q_count_168h, type_a2q_reverse_24h, type_a2q_decay_6h, type_a2q_decay_24h, type_a2q_last, type_a2q_src_out_168h, type_a2q_dst_in_168h, type_c2a_count_4h, type_c2a_count_24h, type_c2a_count_168h, type_c2a_reverse_24h, type_c2a_decay_6h, type_c2a_decay_24h, type_c2a_last, type_c2a_src_out_168h, type_c2a_dst_in_168h, type_c2q_count_4h, type_c2q_count_24h, type_c2q_count_168h, type_c2q_reverse_24h, type_c2q_decay_6h, type_c2q_decay_24h, type_c2q_last, type_c2q_src_out_168h, type_c2q_dst_in_168h"],
        ["유형 다양성과 경로 3개", "type_entropy_30d, type_switches_30d, time_ordered_twohop_30d"],
    ],
    widths=[3.7, 12.5],
    font_size=7.4,
)

doc.add_page_break()
add_heading(doc, "참고 자료", 1)
references = [
    "Stanford Network Analysis Project  Stack Exchange temporal network datasets  https://snap.stanford.edu/data/",
    "Hamilton W Ying Z Leskovec J  Inductive Representation Learning on Large Graphs  2017",
    "Rossi E et al  Temporal Graph Networks for Deep Learning on Dynamic Graphs  2020",
    "Ke G et al  LightGBM A Highly Efficient Gradient Boosting Decision Tree  2017",
    "Davis J Goadrich M  The Relationship Between Precision Recall and ROC Curves  ICML 2006",
    "Saito T Rehmsmeier M  The Precision Recall Plot Is More Informative than the ROC Plot When Evaluating Binary Classifiers on Imbalanced Datasets  PLOS ONE 2015",
    "Roberts D R et al  Cross Validation Strategies for Data with Temporal Spatial Hierarchical or Phylogenetic Structure  Ecography 2017",
    "Lahiri S N  Resampling Methods for Dependent Data  Springer 2003",
]
for item in references:
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(0.65)
    p.paragraph_format.first_line_indent = Cm(-0.45)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.05
    r = p.add_run(item)
    set_run_font(r, size=8.8)

OUTPUT.parent.mkdir(parents=True, exist_ok=True)
doc.save(OUTPUT)
print(OUTPUT)

